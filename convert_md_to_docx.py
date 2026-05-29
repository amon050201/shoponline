"""Convert the course design report from Markdown to Word (.docx)."""
import re
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# --- style setup ---
style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.25

# Configure heading styles
for level in range(1, 4):
    heading_style = doc.styles[f'Heading {level}']
    hf = heading_style.font
    if level == 1:
        hf.size = Pt(18)
        hf.bold = True
    elif level == 2:
        hf.size = Pt(15)
        hf.bold = True
    elif level == 3:
        hf.size = Pt(13)
        hf.bold = True
    hf.name = '黑体'
    heading_style.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    hf.color.rgb = RGBColor(0, 0, 0)


def add_styled_paragraph(text, bold=False, size=None, alignment=None, font_name=None):
    """Add a paragraph with optional styling."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    if size:
        run.font.size = Pt(size)
    if font_name:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if alignment is not None:
        p.alignment = alignment
    return p


def add_table_from_md(lines, start, end):
    """Parse a markdown table and add it to the document."""
    table_lines = lines[start:end]
    rows = []
    for line in table_lines:
        line = line.strip().strip('|')
        cells = [c.strip() for c in line.split('|')]
        rows.append(cells)

    if not rows:
        return

    # Filter out separator row
    data_rows = []
    for row in rows:
        if all(re.match(r'^[-:]+$', c) for c in row):
            continue
        data_rows.append(row)

    if not data_rows:
        return

    num_cols = len(data_rows[0])
    table = doc.add_table(rows=len(data_rows), cols=num_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row_data in enumerate(data_rows):
        for j, cell_text in enumerate(row_data):
            if j < num_cols:
                cell = table.cell(i, j)
                cell.text = ''
                p = cell.paragraphs[0]
                run = p.add_run(cell_text)
                run.font.size = Pt(9)
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                if i == 0:
                    run.bold = True
                    shading = cell._element.get_or_add_tcPr()
                    shd = shading.makeelement(qn('w:shd'), {
                        qn('w:fill'): 'D9D9D9',
                        qn('w:val'): 'clear',
                    })
                    shading.append(shd)

    doc.add_paragraph()


def parse_inline_markdown(paragraph, text):
    """Parse inline markdown (bold, code, links) and add runs to paragraph."""
    pattern = r'(\*\*(.+?)\*\*|`(.+?)`|\[(.+?)\]\((.+?)\))'
    last_end = 0
    for match in re.finditer(pattern, text):
        if match.start() > last_end:
            run = paragraph.add_run(text[last_end:match.start()])
            run.font.size = Pt(11)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

        if match.group(2) is not None:  # **bold**
            run = paragraph.add_run(match.group(2))
            run.bold = True
            run.font.size = Pt(11)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        elif match.group(3) is not None:  # `code`
            run = paragraph.add_run(match.group(3))
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif match.group(4) is not None:  # [text](url)
            run = paragraph.add_run(match.group(4))
            run.font.size = Pt(11)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            run.font.color.rgb = RGBColor(0, 0, 255)
            run.underline = True

        last_end = match.end()

    if last_end < len(text):
        run = paragraph.add_run(text[last_end:])
        run.font.size = Pt(11)
        run.font.name = '宋体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def add_code_block(code_lines):
    """Add a code block with gray background."""
    for code_line in code_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        pPr = p._element.get_or_add_pPr()
        shd = pPr.makeelement(qn('w:shd'), {
            qn('w:fill'): 'F5F5F5',
            qn('w:val'): 'clear',
        })
        pPr.append(shd)
        run = p.add_run(code_line if code_line else ' ')
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    doc.add_paragraph()


def convert_md_to_docx(md_path, docx_path):
    """Main conversion function."""
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line:
            i += 1
            continue

        if line.startswith('---') or line.startswith('==='):
            i += 1
            continue

        if line.startswith('- [') and ']' in line:
            i += 1
            continue

        if line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i].rstrip())
                i += 1
            i += 1
            if code_lines:
                add_code_block(code_lines)
            continue

        heading_match = re.match(r'^(#{1,4})\s+(.+)$', line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            text = re.sub(r'\s*\{#[^}]+\}', '', text)
            doc.add_heading(text, level=level)
            i += 1
            continue

        if '|' in line and i + 1 < len(lines) and re.match(r'^[\|\s\-:]+$', lines[i + 1].strip()):
            start = i
            i += 2
            while i < len(lines) and '|' in lines[i].strip():
                i += 1
            add_table_from_md(lines, start, i)
            continue

        list_match = re.match(r'^(\s*)[-*]\s+(.+)$', line)
        if list_match:
            indent_level = len(list_match.group(1))
            text = list_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0 + indent_level * 0.5)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            bold_match = re.match(r'^\*\*(.+?)\*\*[：:]?\s*(.*)$', text)
            if bold_match:
                run = p.add_run(f"  {bold_match.group(1)}")
                run.bold = True
                run.font.size = Pt(11)
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                if bold_match.group(2):
                    run2 = p.add_run(f"：{bold_match.group(2)}")
                    run2.font.size = Pt(11)
                    run2.font.name = '宋体'
                    run2._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            else:
                parse_inline_markdown(p, text)
            i += 1
            continue

        if line.startswith('> '):
            text = line[2:]
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.5)
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1
            continue

        text = line.strip()
        text = re.sub(r'\s*\{#[^}]+\}', '', text)
        p = doc.add_paragraph()
        parse_inline_markdown(p, text)
        i += 1

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.8)

    doc.save(docx_path)
    print(f"Saved: {docx_path}")


if __name__ == '__main__':
    convert_md_to_docx(
        r'c:\ideaProject\demo1\线上购物与交易系统设计报告.md',
        r'c:\ideaProject\demo1\线上购物与交易系统设计报告.docx'
    )
